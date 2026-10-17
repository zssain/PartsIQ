"""F58 Catalogue imports (owner: worker; depends on 02, 05, 15, F59, F57).

validate_import is a DRY RUN: it reads a staged import file (CSV UTF-8 or XLSX) WITHOUT
evaluating any formula, validates row schema / model references / serial-range overlap /
supersession cycles / version-specific source evidence, and returns a report. It never writes
published data, so an invalid batch changes nothing. Macro-enabled workbooks and formula cells
are rejected outright; a 10 MiB / 50 000-row ceiling is enforced before heavy work. Only
catalogue_admin/admin may validate (and later publish via F59) — a public demo cannot alter the
shared catalogue.

XLSX is read with openpyxl in read_only + data_only=False mode so formulas are seen as formulas
(data_type 'f') and rejected, never computed; macros are detected by inspecting the zip for a
vbaProject part. The live import_batches store is injected and BLOCKED here.
"""
from __future__ import annotations

import csv
import io
import re
import zipfile
from dataclasses import dataclass
from typing import Protocol
from uuid import UUID

from openpyxl import load_workbook

from partsiq_common.contracts import ActorContext, Contract, JobLease
from partsiq_common.errors import PartsIQError

MAX_BYTES = 10 * 1024 * 1024
MAX_ROWS = 50_000
REQUIRED_COLUMNS = frozenset({"part_number", "name_en", "model_id", "source_reference"})
_PUBLISHER_ROLES = frozenset({"catalogue_admin", "admin"})
_SOURCE_RE = re.compile(r"rev\s+\S+\s*,\s*section\s+\S+", re.IGNORECASE)


class ImportErrorView(Contract):
    row: int | None
    code: str
    detail: str


class F58Request(Contract):
    batch_id: UUID
    lease: JobLease


class F58Result(Contract):
    valid: bool
    row_count: int
    errors: list[ImportErrorView]


@dataclass(frozen=True, slots=True)
class ImportFile:
    filename: str
    media_type: str
    data: bytes


class ImportStore(Protocol):
    async def load_batch(self, batch_id: UUID) -> ImportFile | None: ...

    async def known_model_ids(self) -> set[int]: ...


def _authorize(actor: ActorContext) -> None:
    if not (_PUBLISHER_ROLES & set(actor.roles)):
        raise PartsIQError("FORBIDDEN", internal_detail="catalogue import requires catalogue_admin/admin")


def _is_formula(value: object) -> bool:
    # CSV-injection / formula vectors: a leading = or @ (openpyxl also flags data_type 'f').
    return isinstance(value, str) and value[:1] in {"=", "@"}


def _parse_csv(data: bytes) -> tuple[list[dict], list[ImportErrorView]]:
    try:
        text = data.decode("utf-8")
    except UnicodeDecodeError:
        return [], [ImportErrorView(row=None, code="ENCODING", detail="file must be UTF-8")]
    reader = csv.DictReader(io.StringIO(text))
    header = set(reader.fieldnames or [])
    missing = REQUIRED_COLUMNS - header
    if missing:
        return [], [ImportErrorView(row=None, code="SCHEMA", detail=f"missing columns: {sorted(missing)}")]
    rows: list[dict] = []
    errors: list[ImportErrorView] = []
    for line_no, raw in enumerate(reader, start=2):  # header is line 1
        if len(rows) >= MAX_ROWS:
            errors.append(ImportErrorView(row=None, code="TOO_MANY_ROWS", detail=f"more than {MAX_ROWS} rows"))
            break
        if any(_is_formula(v) for v in raw.values()):
            errors.append(ImportErrorView(row=line_no, code="FORMULA_NOT_ALLOWED", detail="formula-like cell rejected"))
            continue
        rows.append({k: (raw.get(k) or "") for k in raw} | {"_row": line_no})
    return rows, errors


def _parse_xlsx(data: bytes) -> tuple[list[dict], list[ImportErrorView]]:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
    except zipfile.BadZipFile:
        return [], [ImportErrorView(row=None, code="MALFORMED", detail="not a valid XLSX file")]
    if any("vbaProject" in name for name in names):
        return [], [ImportErrorView(row=None, code="MACROS_NOT_ALLOWED", detail="macro-enabled workbook rejected")]

    workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False)  # never evaluate formulas
    worksheet = workbook.active
    rows: list[dict] = []
    errors: list[ImportErrorView] = []
    header: list[str] = []
    for index, cells in enumerate(worksheet.iter_rows(), start=1):
        if index == 1:
            header = [str(c.value) if c.value is not None else "" for c in cells]
            missing = REQUIRED_COLUMNS - set(header)
            if missing:
                workbook.close()
                return [], [ImportErrorView(row=None, code="SCHEMA", detail=f"missing columns: {sorted(missing)}")]
            continue
        if len(rows) >= MAX_ROWS:
            errors.append(ImportErrorView(row=None, code="TOO_MANY_ROWS", detail=f"more than {MAX_ROWS} rows"))
            break
        if any(c.data_type == "f" or _is_formula(c.value) for c in cells):
            errors.append(ImportErrorView(row=index, code="FORMULA_NOT_ALLOWED", detail="formula cell rejected"))
            continue
        values = ["" if c.value is None else str(c.value) for c in cells]
        rows.append({header[i]: values[i] if i < len(values) else "" for i in range(len(header))} | {"_row": index})
    workbook.close()
    return rows, errors


def _parse(file: ImportFile) -> tuple[list[dict], list[ImportErrorView]]:
    name = file.filename.lower()
    if name.endswith(".csv") or file.media_type in {"text/csv", "application/csv"}:
        return _parse_csv(file.data)
    if name.endswith(".xlsm"):
        return [], [ImportErrorView(row=None, code="MACROS_NOT_ALLOWED", detail="macro-enabled workbook rejected")]
    if name.endswith(".xlsx") or "spreadsheetml" in file.media_type:
        return _parse_xlsx(file.data)
    return [], [ImportErrorView(row=None, code="UNSUPPORTED_FORMAT", detail="use CSV (UTF-8) or XLSX")]


def _validate_rows(rows: list[dict], known_models: set[int]) -> list[ImportErrorView]:
    errors: list[ImportErrorView] = []
    serial_ranges: dict[int, list[tuple[int, int, int]]] = {}  # model_id -> [(start, end, row)]
    supersession: dict[str, str] = {}
    for row in rows:
        line = row["_row"]
        if not row.get("part_number"):
            errors.append(ImportErrorView(row=line, code="ROW_SCHEMA", detail="part_number is required"))
        try:
            model_id = int(row.get("model_id") or "")
        except ValueError:
            errors.append(ImportErrorView(row=line, code="BAD_REFERENCE", detail="model_id must be an integer"))
            model_id = None  # type: ignore[assignment]
        if model_id is not None and model_id not in known_models:
            errors.append(ImportErrorView(row=line, code="BAD_REFERENCE", detail=f"unknown model_id {model_id}"))
        if not _SOURCE_RE.search(row.get("source_reference") or ""):
            errors.append(ImportErrorView(row=line, code="MISSING_EVIDENCE", detail="source reference not version-specific"))
        start, end = row.get("serial_start") or "", row.get("serial_end") or ""
        if start and end:
            try:
                lo, hi = int(start), int(end)
                serial_ranges.setdefault(model_id, []).append((lo, hi, line))  # type: ignore[arg-type]
            except ValueError:
                errors.append(ImportErrorView(row=line, code="BAD_SERIAL", detail="serial range must be integers"))
        superseded_by = row.get("superseded_by") or ""
        if superseded_by and row.get("part_number"):
            supersession[row["part_number"]] = superseded_by

    for ranges in serial_ranges.values():
        ordered = sorted(ranges)
        for (lo1, hi1, _), (lo2, hi2, row2) in zip(ordered, ordered[1:], strict=False):
            if lo2 <= hi1:
                errors.append(ImportErrorView(row=row2, code="SERIAL_OVERLAP", detail="serial range overlaps another"))

    if _has_cycle(supersession):
        errors.append(ImportErrorView(row=None, code="SUPERSESSION_CYCLE", detail="supersession graph contains a cycle"))
    return errors


def _has_cycle(edges: dict[str, str]) -> bool:
    """Single-successor supersession graph: follow each chain; a revisit means a cycle."""
    for start in edges:
        seen = set()
        node = start
        while node in edges:
            if node in seen:
                return True
            seen.add(node)
            node = edges[node]
    return False


async def validate_import(actor: ActorContext, request: F58Request, *, store: ImportStore) -> F58Result:
    """Dry-run validation of a staged import batch; never writes published data."""
    _authorize(actor)
    file = await store.load_batch(request.batch_id)
    if file is None:
        raise PartsIQError("NOT_FOUND", internal_detail="import batch not found")
    if len(file.data) > MAX_BYTES:
        return F58Result(
            valid=False, row_count=0,
            errors=[ImportErrorView(row=None, code="FILE_TOO_LARGE", detail=f"file exceeds {MAX_BYTES} bytes")],
        )
    rows, parse_errors = _parse(file)
    errors = parse_errors + _validate_rows(rows, await store.known_model_ids())
    return F58Result(valid=not errors, row_count=len(rows), errors=errors)


# --- worker wiring: lease-driven job entry (owner registers the store at startup; BLOCKED here) ---
_VALID_ROLES = frozenset({"reviewer", "catalogue_admin", "admin"})
_store: ImportStore | None = None


def configure_store(store: ImportStore | None) -> None:
    global _store
    _store = store


async def validate_import_job(lease: JobLease) -> F58Result:
    """Dispatcher adapter: re-check admin from the signed payload and run the dry-run validation."""
    from uuid import UUID as _UUID

    if _store is None:
        raise PartsIQError("DEPENDENCY_UNAVAILABLE", internal_detail="import store not configured")
    payload = lease.payload
    roles = [r for r in str(payload.get("actor_roles") or "").split(",") if r in _VALID_ROLES]
    actor = ActorContext(
        user_id=_UUID(str(payload["actor_user_id"])),
        request_id=_UUID(str(payload["request_id"])),
        roles=roles,  # type: ignore[arg-type]
    )
    return await validate_import(actor, F58Request(batch_id=_UUID(str(payload["batch_id"])), lease=lease), store=_store)
